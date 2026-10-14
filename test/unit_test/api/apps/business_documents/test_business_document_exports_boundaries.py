"""Export attempts are fenced before storage and before metadata publication."""

from dataclasses import replace
from io import BytesIO

import pytest
from docx import Document

from test.unit_test.api.apps.business_documents.test_business_document_worker_exports import (
    AUTHOR,
    TENANT,
    MemoryStorage,
    _agreed_document,
    _command,
    _request_export,
    database as worker_database,
)
from api.apps.business_documents.exports import BusinessDocumentExportService
from api.apps.business_documents.runtime import document_commands, job_completion
from api.apps.business_documents.worker import BusinessDocumentJobQueue
from api.db.db_models import BusinessDocument, BusinessDocumentEvent, BusinessDocumentExportArtifact, BusinessDocumentExportStage, BusinessDocumentJob, BusinessDocumentRevision
from business_documents.application.errors import BusinessDocumentError, ConflictError


@pytest.fixture()
def database():
    yield from worker_database.__wrapped__()


def export_attempt():
    document = _agreed_document()
    requested = document_commands.execute(
        TENANT, AUTHOR, document["document_id"], _command(document, "REQUEST_EXPORT", {"revision_id": document["current_revision"]["revision_id"], "format": "MARKDOWN"})
    )
    job = BusinessDocumentJobQueue.claim("boundary-worker", lease_ms=60_000)
    assert job.id == requested["job_id"]
    return document, job


@pytest.mark.parametrize(
    "damage, code",
    [
        ("job_type", "INVALID_EXPORT_JOB"),
        ("not_running", "JOB_LEASE_LOST"),
        ("no_token", "JOB_LEASE_LOST"),
        ("foreign_tenant", "DOCUMENT_NOT_FOUND"),
        ("revision_changed", "AGREED_REVISION_REQUIRED"),
        ("version_changed", "STALE_AI_RESULT"),
        ("unsupported_format", "INVALID_EXPORT_FORMAT"),
        ("missing_revision", "REVISION_NOT_FOUND"),
    ],
)
def test_invalid_export_prerequisite_never_writes_storage(database, damage, code):
    document, job = export_attempt()
    storage = MemoryStorage()
    if damage == "job_type":
        job.job_type = "GENERATE_DRAFT"
    elif damage == "not_running":
        job.status = "CANCELLED"
    elif damage == "no_token":
        job.lease_token = None
    elif damage == "foreign_tenant":
        job.tenant_id = "foreign-tenant"
    elif damage == "revision_changed":
        BusinessDocument.update(current_revision_id="new-revision").where(BusinessDocument.id == document["document_id"]).execute()
    elif damage == "version_changed":
        BusinessDocument.update(state_version=job.source_state_version + 1).where(BusinessDocument.id == document["document_id"]).execute()
    elif damage == "unsupported_format":
        job.payload = {**job.payload, "command_payload": {**job.payload["command_payload"], "format": "PDF"}}
    elif damage == "missing_revision":
        BusinessDocumentRevision.delete().where(BusinessDocumentRevision.id == job.payload["command_payload"]["revision_id"]).execute()
    with pytest.raises(BusinessDocumentError) as caught:
        BusinessDocumentExportService.generate(job, storage=storage)
    assert caught.value.code == code
    assert storage.put_count == 0
    assert BusinessDocumentExportStage.select().count() == 0
    assert BusinessDocumentExportArtifact.select().count() == 0


@pytest.mark.parametrize("cancel_at", [1, 2, 3])
def test_cancellation_during_export_preserves_cleanup_evidence(database, cancel_at):
    _, job = export_attempt()
    storage = MemoryStorage()
    calls = 0

    def ensure_current():
        nonlocal calls
        calls += 1
        if calls == cancel_at:
            raise ConflictError("JOB_LEASE_LOST", "Attempt was cancelled")

    with pytest.raises(BusinessDocumentError) as caught:
        BusinessDocumentExportService.generate(job, storage=storage, ensure_current=ensure_current)
    assert caught.value.code == "JOB_LEASE_LOST"
    assert BusinessDocumentExportArtifact.select().count() == 0
    assert storage.put_count == (1 if cancel_at == 3 else 0)
    if cancel_at == 3:
        stage = BusinessDocumentExportStage.get()
        assert stage.state == "RESERVED"
        assert BusinessDocumentExportService.cleanup_stage(stage.id, storage=storage)
        assert storage.objects == {}
    assert BusinessDocumentExportStage.select().count() == 0


@pytest.mark.parametrize("damage", ["removed", "cleanup_claimed", "different_bytes"])
def test_changed_stage_blocks_atomic_publication(database, damage):
    _, job = export_attempt()
    storage = MemoryStorage()
    prepared = BusinessDocumentExportService.generate(job, storage=storage)
    if damage == "removed":
        BusinessDocumentExportStage.delete().where(BusinessDocumentExportStage.id == prepared.stage_id).execute()
    else:
        changes = {"state": "CLEANING"} if damage == "cleanup_claimed" else {"content_hash": "sha256:" + "0" * 64}
        BusinessDocumentExportStage.update(**changes).where(BusinessDocumentExportStage.id == prepared.stage_id).execute()
    events_before = BusinessDocumentEvent.select().count()
    with pytest.raises(BusinessDocumentError) as caught:
        job_completion.complete(TENANT, job.lease_owner, job.id, prepared, job.lease_token)
    assert caught.value.code == "EXPORT_STAGE_CHANGED"
    assert BusinessDocumentExportArtifact.select().count() == 0
    assert BusinessDocumentEvent.select().count() == events_before
    assert BusinessDocumentJob.get_by_id(job.id).status == "RUNNING"


@pytest.mark.parametrize(
    "changes",
    [
        {"tenant_id": "foreign-tenant"},
        {"lease_token": "previous-attempt"},
        {"storage_key": "exports/another-document/private.md"},
        {"content_hash": "unverified"},
    ],
)
def test_prepared_export_identity_is_not_trusted(database, changes):
    _, job = export_attempt()
    prepared = BusinessDocumentExportService.generate(job, storage=MemoryStorage())
    with pytest.raises(BusinessDocumentError) as caught:
        job_completion.complete(TENANT, job.lease_owner, job.id, replace(prepared, **changes), job.lease_token)
    assert caught.value.code == "EXPORT_JOB_CHANGED"
    assert BusinessDocumentExportArtifact.select().count() == 0
    assert BusinessDocumentExportStage.get_by_id(prepared.stage_id).state == "STORED"
    assert BusinessDocumentJob.get_by_id(job.id).status == "RUNNING"


def test_reused_artifact_cannot_be_published_after_concurrent_metadata_change(database):
    storage = MemoryStorage()
    document = _agreed_document()
    artifact, document = _request_export(document, storage, "MARKDOWN")
    document_commands.execute(TENANT, AUTHOR, document["document_id"], _command(document, "REQUEST_EXPORT", {"revision_id": artifact["revision_id"], "format": "MARKDOWN"}))
    job = BusinessDocumentJobQueue.claim("boundary-worker", lease_ms=60_000)
    prepared = BusinessDocumentExportService.generate(job, storage=storage)
    assert prepared.created_blob is False
    BusinessDocumentExportArtifact.update(content_hash="sha256:" + "0" * 64).where(BusinessDocumentExportArtifact.id == artifact["artifact_id"]).execute()
    with pytest.raises(BusinessDocumentError) as caught:
        job_completion.complete(TENANT, job.lease_owner, job.id, prepared, job.lease_token)
    assert caught.value.code == "EXPORT_ARTIFACT_CHANGED"
    assert BusinessDocumentExportArtifact.select().count() == 1
    assert BusinessDocumentJob.get_by_id(job.id).status == "RUNNING"


def test_download_translates_storage_failure_without_returning_unverified_bytes(database):
    storage = MemoryStorage()
    document = _agreed_document()
    artifact, _ = _request_export(document, storage, "MARKDOWN")

    class UnavailableStorage:
        def get(self, *_args):
            raise OSError("Object store unavailable")

    with pytest.raises(BusinessDocumentError) as caught:
        BusinessDocumentExportService.download(TENANT, AUTHOR, document["document_id"], artifact["artifact_id"], storage=UnavailableStorage())
    assert caught.value.code == "EXPORT_STORAGE_CORRUPT"
    assert caught.value.status == 500
    assert isinstance(caught.value.__cause__, OSError)


@pytest.mark.parametrize("missing", ["document", "artifact"])
def test_download_missing_identity_does_not_access_object_store(database, missing):
    document = _agreed_document()
    storage = MemoryStorage()
    document_id = "absent-document" if missing == "document" else document["document_id"]
    with pytest.raises(BusinessDocumentError) as caught:
        BusinessDocumentExportService.download(TENANT, AUTHOR, document_id, "absent-artifact", storage=storage)
    assert caught.value.code == ("DOCUMENT_NOT_FOUND" if missing == "document" else "EXPORT_NOT_FOUND")
    assert caught.value.status == 404
    assert storage.objects == {}


def test_cleanup_of_an_absent_stage_is_idempotent(database):
    storage = MemoryStorage()
    storage.put("other-run", "keep", b"unrelated")
    assert BusinessDocumentExportService.cleanup_stage("already-cleaned-stage", storage=storage)
    assert BusinessDocumentExportService.cleanup_stage("already-cleaned-stage", storage=storage)
    assert storage.objects == {("other-run", "keep"): b"unrelated"}


def test_ambiguous_cleanup_acknowledgment_retains_blob_ledger(database):
    _, job = export_attempt()

    class AmbiguousStorage(MemoryStorage):
        def remove_and_confirm_absent(self, _bucket, _key):
            return 1

    storage = AmbiguousStorage()
    prepared = BusinessDocumentExportService.generate(job, storage=storage)
    assert BusinessDocumentExportService.cleanup_stage(prepared.stage_id, storage=storage) is False
    stage = BusinessDocumentExportStage.get_by_id(prepared.stage_id)
    assert stage.state == "CLEANING"
    assert stage.cleanup_after > 0
    assert storage.get(prepared.storage_bucket, prepared.storage_key) is not None
    assert BusinessDocumentExportArtifact.select().count() == 0


def test_docx_contains_lists_diagrams_links_and_rectangular_tables():
    ast = {
        "sections": [
            {
                "id": "3.1",
                "title": "Scope",
                "blocks": [
                    {"type": "paragraph", "text": "Introduction"},
                    {"type": "list", "items": ["First", "Second"]},
                    {"type": "table", "headers": ["Name", "Value"], "rows": [["empty", None], ["bounded", "value", "outside-table"]]},
                    {"type": "plantuml", "source": "@startuml\nA -> B\n@enduml"},
                    {"type": "image", "alt": "Diagram", "url": "https://example.invalid/diagram.png"},
                    {"type": "reference", "label": "Reference", "url": "https://example.invalid/reference"},
                ],
            }
        ]
    }
    document = Document(BytesIO(BusinessDocumentExportService._render_docx(ast)))
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    assert paragraphs == ["3.1. Scope", "Introduction", "First", "Second", "@startuml\nA -> B\n@enduml", "Diagram: https://example.invalid/diagram.png", "Reference: https://example.invalid/reference"]
    assert [paragraph.style.name for paragraph in document.paragraphs[2:4]] == ["List Bullet", "List Bullet"]
    assert [[cell.text for cell in row.cells] for row in document.tables[0].rows] == [["Name", "Value"], ["empty", ""], ["bounded", "value"]]
    assert "outside-table" not in " ".join(cell.text for row in document.tables[0].rows for cell in row.cells)


@pytest.mark.parametrize("kind", ["image", "reference"])
@pytest.mark.parametrize("url", ["javascript:alert(1)", "file:///private-document", "/relative/path"])
def test_eva_export_rejects_links_without_absolute_http_origin(kind, url):
    block = {"type": kind, "url": url, **({"alt": "Image"} if kind == "image" else {"label": "Reference"})}
    with pytest.raises(BusinessDocumentError) as caught:
        BusinessDocumentExportService._render_eva_wiki({"sections": [{"id": "1", "title": "Scope", "blocks": [block]}]}, 1)
    assert caught.value.code == "UNSAFE_EXPORT_URL"


def test_eva_export_drops_empty_paragraphs_and_list_items():
    ast = {
        "sections": [
            {
                "id": "1",
                "title": "Scope",
                "blocks": [
                    {"type": "paragraph", "text": "  "},
                    {"type": "list", "items": [" ", ""]},
                    {"type": "list", "items": ["", "A < B"]},
                ],
            }
        ]
    }
    html = BusinessDocumentExportService._render_eva_wiki(ast, 1)
    assert html.count("<ul ") == 1 and html.count("<li ") == 1
    assert "A &lt; B" in html
    assert 'data-id="br-r1-s1-b1"' not in html
    assert 'data-id="br-r1-s1-b2"' not in html
